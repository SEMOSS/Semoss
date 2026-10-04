"""Real local conversions: SEMOSS_TEST_DOCLING=1 pytest this file.

The first PDF conversion can download Docling's model artifacts. Fixtures are
generated locally and do not contain user documents or require a model endpoint.
"""

import base64
import os
from io import BytesIO

import pytest
from genai_client.message_builders.openai.openai_message_builder import (
    OpenAIMessageBuilder,
)
from genai_client.message_builders.semoss_base.document_input import (
    DocumentInputProcessor,
)
from genai_client.message_builders.semoss_base.semoss_message_builder import (
    SEMOSSMessageBuilder,
)
from genai_client.message_builders.semoss_base.semoss_models import ModelSettings

pytestmark = pytest.mark.skipif(
    os.environ.get("SEMOSS_TEST_DOCLING") != "1",
    reason="Set SEMOSS_TEST_DOCLING=1 to run local Docling conversions",
)


@pytest.fixture(scope="module")
def processor():
    return DocumentInputProcessor()


def make_document(extension):
    output = BytesIO()
    if extension == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Quarterly report", 0)
        document.add_paragraph("Revenue is 42 million.")
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [("Region", "Revenue"), ("East", "42")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.save(output)
    elif extension == "pptx":
        from pptx import Presentation

        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "Quarterly report"
        slide.placeholders[1].text = "Revenue is 42 million."
        slide.notes_slide.notes_text_frame.text = "Launch in October."
        document.save(output)
    elif extension == "xlsx":
        from openpyxl import Workbook

        document = Workbook()
        sheet = document.active
        sheet.title = "Revenue"
        sheet.append(["Region", "Revenue"])
        sheet.append(["East", 42])
        other = document.create_sheet("Expenses")
        other.append(["Category", "Amount"])
        other.append(["Travel", 7])
        document.save(output)
    elif extension == "pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        document = PdfWriter()
        page = document.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): document._add_object(font)}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(
            b"BT /F1 18 Tf 72 720 Td (Quarterly report) Tj 0 -30 Td (Revenue is 42 million.) Tj ET"
        )
        page[NameObject("/Contents")] = document._add_object(stream)
        document.write(output)
    return output.getvalue()


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "pdf"])
def test_real_document_to_text_payload(processor, extension):
    raw = make_document(extension)
    source = [
        {
            "type": "INPUT_TEXT",
            "io": "INPUT",
            "schemaVersion": 2,
            "messageId": f"fixture-{extension}",
            "parts": [
                {"type": "TEXT", "text": "What is the revenue?"},
                {
                    "type": "MEDIA",
                    "mediaInfo": {
                        "fileName": f"report.{extension}",
                        "base64Data": base64.b64encode(raw).decode(),
                    },
                },
            ],
        }
    ]
    settings = ModelSettings(
        model_name="fixture",
        context_window=32768,
        max_tokens=1024,
    )
    messages = SEMOSSMessageBuilder().build_messages(source, {}, settings)
    prepared = processor.prepare(
        messages, settings, source_messages=source, input_modalities=["TEXT"]
    )
    wire = OpenAIMessageBuilder(settings, "chat-completion").build_request(prepared)
    text = wire["messages"][0]["content"][1]["text"]
    assert "42" in text
    assert (
        messages[0].parts[1].media_info.data
        == source[0]["parts"][1]["mediaInfo"]["base64Data"]
    )
    if extension == "pptx":
        assert "Slide 1" in text
        assert "Launch in October" in text
    elif extension == "xlsx":
        assert "Revenue" in text and "Expenses" in text
        assert "Travel" in text and "East" in text
    elif extension == "pdf":
        assert "Page 1" in text
    elif extension == "docx":
        assert "East" in text and "|" in text
